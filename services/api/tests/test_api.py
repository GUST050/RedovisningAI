"""API-tester mot riktig databas (applikationsrollen) med utvecklingsinloggning."""

from __future__ import annotations

import io
import json
import uuid
import zipfile
from copy import deepcopy
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from redovisningai.ai.service import AIService, FakeProvider
from redovisningai.api import routes_other
from redovisningai.api.app import create_app
from redovisningai.db import models as m
from redovisningai.db.bootstrap import add_user, create_company, create_organization
from redovisningai.db.session import TenantContext, tenant_session
from redovisningai.devdata.generator import DEMO_PROFILES, generate
from redovisningai.jobs.pipeline import import_sie
from redovisningai.sie.writer import write_sie4

pytestmark = pytest.mark.usefixtures("database")


@pytest.fixture(scope="module")
def env(database):  # type: ignore[no-untyped-def]
    org = create_organization("API-byrån", "admin@api.se", "Admin", owner_url=database)
    kalle = add_user(org.org_id, "kalle@api.se", "Kalle", "CONSULTANT", owner_url=database)
    add_user(org.org_id, "vera@api.se", "Vera", "VIEWER", owner_url=database)
    admin = TenantContext(org.org_id, org.admin_user_id, "ADMIN", True, True, "admin@api.se")
    p = DEMO_PROFILES[0]
    cid = create_company(admin, p.name, "556677-8800", assign_to=[kalle], vat_period=p.vat_period)
    g = generate(p, date(2026, 10, 12))
    for y in g.ledger.years:
        raw = write_sie4(g.ledger, y).replace(b"#ORGNR 556677-8899", b"#ORGNR 556677-8800")
        import_sie(admin, cid, f"{y.fiscal_year.label}.se", raw)
    client = TestClient(create_app())
    return {"client": client, "cid": str(cid), "org": org}


def H(email: str) -> dict[str, str]:
    return {"X-Dev-User": email}


def test_requires_login(env) -> None:  # type: ignore[no-untyped-def]
    assert env["client"].get("/api/me").status_code == 401
    assert env["client"].get("/api/me", headers=H("okand@x.se")).status_code == 403


def test_encoded_dev_login_cookie_authenticates(env) -> None:  # type: ignore[no-untyped-def]
    response = env["client"].get("/api/me", headers={"Cookie": "rai_dev_user=admin%40api.se"})
    assert response.status_code == 200
    assert response.json()["user"]["email"] == "admin@api.se"


def test_me_and_security_headers(env) -> None:  # type: ignore[no-untyped-def]
    r = env["client"].get("/api/me", headers=H("admin@api.se"))
    assert r.status_code == 200
    assert r.json()["role"] == "ADMIN"
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert "default-src 'none'" in r.headers["Content-Security-Policy"]


def test_portfolio_prioritizes_and_explains(env) -> None:  # type: ignore[no-untyped-def]
    r = env["client"].get("/api/portfolio", headers=H("kalle@api.se")).json()
    assert len(r["companies"]) == 1
    c = r["companies"][0]
    assert c["priority"]["score"] > 0 and c["priority"]["reasons"]
    assert c["open_findings"]["high"] > 0


def test_overview_statements_explain(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    ov = cl.get(f"/api/companies/{cid}/overview?period=2026-09", headers=H("kalle@api.se")).json()
    assert ov["sections"]["ytd"]["metrics"]["net_sales"]["fact"]["display"].endswith("Mkr")
    st = cl.get(f"/api/companies/{cid}/statements?period=YTD:2026-09", headers=H("kalle@api.se")).json()
    assert st["income"]["lines"][0]["code"] == "net_sales"
    ex = cl.get(
        f"/api/companies/{cid}/explain?target=operating_result&period=YTD:2026-09", headers=H("kalle@api.se")
    ).json()
    assert ex["bridge"]["components"]
    dd = cl.get(
        f"/api/companies/{cid}/explain?target=account:6550&period=YTD:2026-09", headers=H("kalle@api.se")
    ).json()
    assert dd["drilldown"]["counterparties"][0]["is_new"]
    assert cl.get(f"/api/companies/{cid}/explain?target=bogus", headers=H("kalle@api.se")).status_code == 422


def test_payroll_masking(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    tx = cl.get(f"/api/companies/{cid}/transactions?account=7210", headers=H("kalle@api.se")).json()
    assert tx["total"] == 0  # RLS: lönerader dolda
    tx_admin = cl.get(f"/api/companies/{cid}/transactions?account=7210", headers=H("admin@api.se")).json()
    assert tx_admin["total"] > 0
    v = cl.get(f"/api/companies/{cid}/vouchers/L1", headers=H("kalle@api.se")).json()
    assert v.get("payroll_rows_masked") and all(not 7000 <= r["account"] <= 7699 for r in v["rows"])


def test_other_client_not_visible(env) -> None:  # type: ignore[no-untyped-def]
    cl = env["client"]
    other = create_company(
        TenantContext(env["org"].org_id, env["org"].admin_user_id, "ADMIN"), "Hemlig AB", "556000-1111"
    )
    assert cl.get(f"/api/companies/{other}/overview", headers=H("kalle@api.se")).status_code == 404
    assert cl.get(f"/api/companies/{other}", headers=H("kalle@api.se")).status_code == 404


def test_cases_decide_and_viewer_blocked(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    cases = cl.get(f"/api/companies/{cid}/cases?period=2026-09", headers=H("kalle@api.se")).json()
    assert cases and all(c["visibility"] != "RESTRICTED_AML" for c in cases)
    assert all("{f:" not in c["root_cause"] for c in cases)
    low = next(c for c in cases if c["severity"] != "HIGH")
    r = cl.post(
        f"/api/companies/{cid}/cases/{low['key']}/decision",
        headers=H("vera@api.se"),
        json={"status": "ACCEPTED_OK", "note": "ok"},
    )
    assert r.status_code == 403
    r = cl.post(
        f"/api/companies/{cid}/cases/{low['key']}/decision", headers=H("kalle@api.se"), json={"status": "ACCEPTED_OK"}
    )
    assert r.status_code == 422  # motivering krävs
    r = cl.post(
        f"/api/companies/{cid}/cases/{low['key']}/decision",
        headers=H("kalle@api.se"),
        json={"status": "ACCEPTED_OK", "note": "Känt mönster"},
    )
    assert r.status_code == 200 and r.json()["updated"] >= 1
    mem = cl.get(f"/api/companies/{cid}/memory", headers=H("kalle@api.se")).json()
    assert any(x["rationale"] == "Känt mönster" for x in mem)


def test_question_roundtrip(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    cases = cl.get(f"/api/companies/{cid}/cases?period=2026-09", headers=H("kalle@api.se")).json()
    case = next(c for c in cases if c["status"] != "CLOSED")
    draft = cl.get(f"/api/companies/{cid}/cases/{case['key']}/question-draft", headers=H("kalle@api.se")).json()
    assert draft["text"]
    r = cl.post(
        f"/api/companies/{cid}/questions",
        headers=H("kalle@api.se"),
        json={"case_key": case["key"], "text": draft["text"]},
    ).json()
    token = r["link"].rsplit("/", 1)[1]
    pub = cl.get(f"/api/public/questions/{token}")
    assert pub.status_code == 200 and pub.json()["company_name"] == "Bygg & Co AB"
    ans = cl.post(
        f"/api/public/questions/{token}",
        data={"answer": "Förskott till leverantör"},
        files=[("files", ("kvitto.pdf", b"%PDF-1.4 test", "application/pdf"))],
    )
    assert ans.status_code == 200 and ans.json()["status"] == "ANSWERED"
    bad = cl.post(
        f"/api/public/questions/{token}",
        data={"answer": "x"},
        files=[("files", ("virus.exe", b"MZ", "application/x-msdownload"))],
    )
    assert bad.status_code == 422
    assert cl.get("/api/public/questions/ogiltig-token-123456789").status_code == 404
    qs = cl.get(f"/api/companies/{cid}/questions", headers=H("kalle@api.se")).json()
    att = qs[0]["attachments"][0]
    dl = cl.get(f"/api/companies/{cid}/questions/{qs[0]['id']}/attachments/{att['key']}", headers=H("kalle@api.se"))
    assert dl.content == b"%PDF-1.4 test"


def test_ai_endpoints_fall_back_without_provider(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    c = cl.post(f"/api/companies/{cid}/periods/2026-09/commentary", headers=H("kalle@api.se")).json()
    assert c["source"] == "rules" and c["data"]["claims"]
    assert all("{f:" not in x["rendered"] for x in c["data"]["claims"])
    m = cl.post(f"/api/companies/{cid}/periods/2026-09/meeting", headers=H("kalle@api.se")).json()
    assert "summary" in m["data"]
    detail = cl.get(f"/api/companies/{cid}/periods/2026-09", headers=H("kalle@api.se")).json()
    assert detail["commentary"]["stale"] is False
    assert detail["client_report"]["stale"] is False
    a = cl.post(f"/api/companies/{cid}/ask", headers=H("kalle@api.se"), json={"question": "Varför föll resultatet?"})
    assert a.status_code == 200


def test_fresh_meeting_can_be_approved_and_exported(env) -> None:  # type: ignore[no-untyped-def]
    client, company_id = env["client"], env["cid"]
    created = client.post(f"/api/companies/{company_id}/periods/2026-09/meeting", headers=H("kalle@api.se"))
    assert created.status_code == 200
    statement = "Försäljningen ska följas upp med kunden."
    approved = client.put(
        f"/api/companies/{company_id}/periods/2026-09/meeting",
        headers=H("admin@api.se"),
        json={
            "summary": [statement],
            "questions": [],
            "decisions": [
                {"index": 0, "statement": statement, "decision": "approve", "reason": "Kontrollerat mot bokföringen"}
            ],
            "approve": True,
        },
    )
    assert approved.status_code == 200
    exported = client.get(
        f"/api/companies/{company_id}/reports/client?period=2026-09&format=docx",
        headers=H("admin@api.se"),
    )
    assert exported.status_code == 200


def test_a3_sends_only_aggregated_transaction_findings_to_provider(env, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    provider = FakeProvider()
    service = AIService(provider, keep_payloads=False)
    monkeypatch.setattr("redovisningai.api.routes_review._ai", lambda _principal: service)

    response = cl.post(f"/api/companies/{cid}/periods/2026-09/commentary", headers=H("kalle@api.se"))

    assert response.status_code == 200
    assert len(provider.calls) == 1
    user_content = provider.calls[0]["user_content"]
    start = user_content.index("<kunddata>") + len("<kunddata>\n")
    end = user_content.index("\n</kunddata>")
    package = json.loads(user_content[start:end])
    assert "company" not in package
    assert len(package["findings"]) <= 5
    assert all("sources" not in finding for finding in package["findings"])
    assert all(finding["label"] != "private vendor name" for finding in package["findings"])
    assert all("voucher" not in fact and "lineage" not in fact and "label" not in fact for fact in package["facts"])
    serialized = json.dumps(package, ensure_ascii=False)
    assert "API-byrån" not in serialized
    assert "Kalle" not in serialized
    assert "lönespecifikation" not in serialized
    assert response.json()["trace_id"]


def test_reports_and_exports(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    pdf = cl.get(f"/api/companies/{cid}/reports/client?period=2026-09&format=pdf", headers=H("kalle@api.se"))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    docx = cl.get(f"/api/companies/{cid}/reports/client?period=2026-09&format=docx", headers=H("kalle@api.se"))
    assert docx.content[:2] == b"PK"
    xlsx = cl.get(f"/api/companies/{cid}/export/statements.xlsx?period=YTD:2026-09", headers=H("kalle@api.se"))
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    reko = cl.get(f"/api/companies/{cid}/reports/reko?period=2026-09", headers=H("kalle@api.se"))
    assert reko.status_code == 200 and reko.content.startswith(b"%PDF")
    # Kundrapporten får inte innehålla interna fynd eller PTL.
    with zipfile.ZipFile(io.BytesIO(docx.content)) as z:
        body = z.read("word/document.xml").decode()
    assert "PTL" not in body and "låneförbud" not in body.lower()


def test_transaction_export_contains_every_visible_row(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    listed = cl.get(
        f"/api/companies/{cid}/transactions?from=2026-01-01&to=2026-09-30&page_size=1",
        headers=H("kalle@api.se"),
    )
    assert listed.status_code == 200
    total = listed.json()["total"]
    assert total > 500  # regression: the export used to stop after its first page

    exported = cl.get(f"/api/companies/{cid}/export/transactions.xlsx?period=YTD:2026-09", headers=H("kalle@api.se"))
    assert exported.status_code == 200
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    try:
        assert workbook["Transaktioner"].max_row - 1 == total
    finally:
        workbook.close()


def test_approve_requires_override_with_open_high(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    r = cl.post(f"/api/companies/{cid}/periods/2026-09/approve", headers=H("kalle@api.se"), json={})
    assert r.status_code == 409
    r = cl.post(
        f"/api/companies/{cid}/periods/2026-09/approve",
        headers=H("kalle@api.se"),
        json={"override_note": "Genomgånget, kunden rättar i oktober"},
    )
    assert r.status_code == 200 and r.json()["status"] == "APPROVED"


def test_aml_restricted(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    assert cl.get(f"/api/companies/{cid}/aml", headers=H("kalle@api.se")).status_code == 403
    r = cl.get(f"/api/companies/{cid}/aml", headers=H("admin@api.se")).json()
    assert r["signals"] and "Meddelandeförbud" in r["notice"]
    sup = cl.post(
        "/api/suppressions",
        headers=H("admin@api.se"),
        json={"rule_code": "AML_LARGE_CASH", "reason": "test", "expires_at": "2027-01-01"},
    )
    assert sup.status_code == 409


def test_rules_catalog_and_audit(env) -> None:  # type: ignore[no-untyped-def]
    cl = env["client"]
    rules = cl.get("/api/rules", headers=H("kalle@api.se")).json()
    assert any(r["code"] == "RELATED_PARTY_RECEIVABLE" and "ABL 21" in r["legal_basis"] for r in rules["rules"])
    assert any(r["code"] == "vat_food" and r["value"] == "0.06" for r in rules["rates"])
    audit = cl.get("/api/audit", headers=H("admin@api.se")).json()
    assert any(e["action"] == "import.completed" for e in audit)


def test_bulk_import_matches_by_orgnr(env) -> None:  # type: ignore[no-untyped-def]
    cl = env["client"]
    g = generate(DEMO_PROFILES[1], date(2026, 10, 12))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("okand.se", write_sie4(g.ledger, g.ledger.current))
    r = cl.post(
        "/api/imports/bulk", headers=H("admin@api.se"), files={"file": ("filer.zip", buf.getvalue(), "application/zip")}
    ).json()
    assert r["results"][0]["status"] == "no_match"


def test_trend_and_voucher_lookup_per_fiscal_year(env) -> None:  # type: ignore[no-untyped-def]
    c, cid = env["client"], env["cid"]
    t = c.get(f"/api/companies/{cid}/trend?months=12", headers=H("kalle@api.se")).json()
    assert len(t["months"]) == 12 and t["months"][-1] == "2026-10"
    assert len(t["net_sales"]) == len(t["costs"]) == len(t["operating_result"]) == 12
    # Verifikationsnummer börjar om varje år: A1 finns både 2025 och 2026.
    v25 = c.get(f"/api/companies/{cid}/vouchers/A1?period=2025-03", headers=H("kalle@api.se")).json()
    v26 = c.get(f"/api/companies/{cid}/vouchers/A1?on=2026-03-01", headers=H("kalle@api.se")).json()
    assert v25["date"].startswith("2025") and v26["date"].startswith("2026")
    latest = c.get(f"/api/companies/{cid}/vouchers/A1", headers=H("kalle@api.se")).json()
    assert latest["date"].startswith("2026")


def test_period_detail_and_unapproved_meeting_not_in_client_report(env) -> None:  # type: ignore[no-untyped-def]
    c, cid = env["client"], env["cid"]
    c.post(f"/api/companies/{cid}/periods/2026-09/meeting", headers=H("kalle@api.se"))
    d = c.get(f"/api/companies/{cid}/periods/2026-09", headers=H("kalle@api.se")).json()
    assert d["client_report"]["approved"] is False
    c.put(
        f"/api/companies/{cid}/periods/2026-09/meeting",
        json={"summary": ["HEMLIGT_UTKAST"], "questions": [], "approve": False},
        headers=H("kalle@api.se"),
    )
    r = c.get(f"/api/companies/{cid}/reports/client?period=2026-09&format=docx", headers=H("kalle@api.se"))
    assert r.status_code == 200
    import docx  # type: ignore[import-untyped]

    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "HEMLIGT_UTKAST" not in text
    assert c.get(f"/api/companies/{cid}/periods/1999-01", headers=H("kalle@api.se")).status_code == 404


def test_client_export_rejects_approved_legacy_draft_without_fingerprint(env) -> None:  # type: ignore[no-untyped-def]
    client, company_id = env["client"], uuid.UUID(env["cid"])
    client.post(f"/api/companies/{company_id}/periods/2026-09/meeting", headers=H("kalle@api.se"))
    org = env["org"]
    ctx = TenantContext(org.org_id, org.admin_user_id, "ADMIN", True, True, "admin@api.se")
    with tenant_session(ctx) as session:
        review = session.scalar(
            select(m.PeriodReview).where(m.PeriodReview.company_id == company_id, m.PeriodReview.period == "2026-09")
        )
        assert review is not None
        original = deepcopy(review.client_report)
        review.client_report = {"approved": True, "data": {"summary": [], "questions": []}}
    try:
        response = client.get(
            f"/api/companies/{company_id}/reports/client?period=2026-09&format=docx",
            headers=H("kalle@api.se"),
        )
        assert response.status_code == 409
    finally:
        with tenant_session(ctx) as session:
            review = session.scalar(
                select(m.PeriodReview).where(
                    m.PeriodReview.company_id == company_id, m.PeriodReview.period == "2026-09"
                )
            )
            assert review is not None
            review.client_report = original


def test_internal_report_uses_saved_commentary_comparison(env, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    client, company_id = env["client"], env["cid"]
    response = client.post(
        f"/api/companies/{company_id}/periods/2026-09/commentary?compare=2026-08",
        headers=H("kalle@api.se"),
    )
    assert response.status_code == 200
    captured: dict[str, str] = {}
    original = routes_other.internal_report

    def capture(overview, *args, **kwargs):  # type: ignore[no-untyped-def]
        captured["compare"] = overview["sections"]["month"]["compare"]["spec"]
        return original(overview, *args, **kwargs)

    monkeypatch.setattr(routes_other, "internal_report", capture)
    exported = client.get(
        f"/api/companies/{company_id}/reports/internal?period=2026-09&format=docx",
        headers=H("kalle@api.se"),
    )
    assert exported.status_code == 200
    assert captured["compare"] == "2026-08"


def test_production_refuses_insecure_defaults() -> None:
    from redovisningai.config import DEV_MASTER_KEY, Settings

    bad = Settings(
        env="prod",
        auth_mode="dev",
        master_key=DEV_MASTER_KEY,
        app_db_password="app",
        database_url="postgresql+psycopg://redovisningai_app:app@db/rai",
    )
    assert len(bad.production_problems()) >= 2
    good = Settings(
        env="prod",
        auth_mode="oidc",
        oidc_issuer="https://login.example",
        oidc_audience="rai",
        oidc_jwks_url="https://login.example/jwks",
        master_key="x" * 48,
        app_db_password="s3cret-long",
        database_url="postgresql+psycopg://redovisningai_app:s3cret-long@db/rai",
    )
    assert good.production_problems() == []


def test_dev_mode_logs_in_default_user_without_login_step(env, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from redovisningai.config import get_settings

    monkeypatch.setattr(get_settings(), "dev_default_user", "admin@api.se")
    r = env["client"].get("/api/me")
    assert r.status_code == 200 and r.json()["user"]["email"] == "admin@api.se"
    # Ett uttryckligt val går före standardanvändaren.
    assert env["client"].get("/api/me", headers=H("kalle@api.se")).json()["user"]["email"] == "kalle@api.se"
    monkeypatch.setattr(get_settings(), "auth_mode", "oidc")
    assert env["client"].get("/api/me").status_code == 401


def test_report_items_rank_and_respect_audience_and_payroll(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    r = cl.get(f"/api/companies/{cid}/report-items?period=YTD:2026-09&mode=yoy", headers=H("kalle@api.se"))
    assert r.status_code == 200
    body = r.json()
    assert body["periods"] == {"current": "YTD:2026-09", "previous": "YTD:2025-09"}
    assert 1 <= len(body["recommended"]) <= 6
    ids = {i["id"] for i in body["items"]}
    assert "metric:net_sales" in ids and set(body["recommended"]) <= ids
    assert {s["kind"] for s in body["series"]} >= {"months", "fiscal_years", "same_month"}
    for item in body["items"]:
        assert not any(f"konto {a}" in item["summary"] for a in range(7000, 7700))
    client = cl.get(
        f"/api/companies/{cid}/report-items?period=2026-09&compare=2026-06&audience=client", headers=H("kalle@api.se")
    ).json()
    assert client["periods"]["previous"] == "2026-06"
    assert all(i["kind"] != "finding" and i["audience"] == "client" for i in client["items"])
    assert (
        cl.get(
            f"/api/companies/{cid}/report-items?period=2026-09&compare=2026-09", headers=H("kalle@api.se")
        ).status_code
        == 422
    )


def test_comparison_report_downloads_in_all_formats(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    items = cl.get(f"/api/companies/{cid}/report-items?period=YTD:2026-09", headers=H("kalle@api.se")).json()
    chosen = [{"id": i, "comment": "Följ upp"} for i in items["recommended"][:3]]
    chosen.append({"id": "structure:operating_margin", "series": "fiscal_years", "count": 3})
    base = {"period": "YTD:2026-09", "mode": "yoy", "items": chosen, "title": "Kvartalsgenomgång"}
    pdf = cl.post(f"/api/companies/{cid}/reports/comparison", headers=H("kalle@api.se"), json={**base, "format": "pdf"})
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
    assert "filename*=UTF-8''" in pdf.headers["content-disposition"]
    # Läsaren Vera är inte tilldelad kunden och ser den inte alls (RLS).
    assert cl.post(f"/api/companies/{cid}/reports/comparison", headers=H("vera@api.se"), json=base).status_code == 404
    docx = cl.post(
        f"/api/companies/{cid}/reports/comparison", headers=H("admin@api.se"), json={**base, "format": "docx"}
    )
    assert docx.status_code == 200
    xlsx = cl.post(
        f"/api/companies/{cid}/reports/comparison",
        headers=H("kalle@api.se"),
        json={**base, "format": "xlsx", "audience": "client"},
    )
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames[0] == "Sammanfattning"
    with tenant_session(TenantContext(env["org"].org_id, env["org"].admin_user_id, "ADMIN")) as s:
        events = s.scalars(select(m.AuditEvent).where(m.AuditEvent.action == "report.comparison_exported")).all()
        assert len(events) >= 3
        assert events[-1].details["audience"] == "client"


def test_comparison_report_rejects_invalid_selection(env) -> None:  # type: ignore[no-untyped-def]
    cl, cid = env["client"], env["cid"]
    url = f"/api/companies/{cid}/reports/comparison"
    unknown = cl.post(url, headers=H("kalle@api.se"), json={"period": "2026-09", "items": [{"id": "metric:okand"}]})
    assert unknown.status_code == 422
    empty = cl.post(url, headers=H("kalle@api.se"), json={"period": "2026-09", "items": []})
    assert empty.status_code == 422
    internal = cl.get(f"/api/companies/{cid}/report-items?period=2026-09", headers=H("admin@api.se")).json()
    finding = next((i["id"] for i in internal["items"] if i["kind"] == "finding"), None)
    if finding is not None:
        client = cl.post(
            url, headers=H("admin@api.se"), json={"period": "2026-09", "audience": "client", "items": [{"id": finding}]}
        )
        assert client.status_code == 422
