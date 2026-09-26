"""Viktigaste skillnaderna och jämförelserapporten (utan databas)."""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from docx import Document as Docx
from openpyxl import load_workbook

from redovisningai.accounting.balances import LedgerIndex
from redovisningai.accounting.categories import CategoryMapping
from redovisningai.accounting.comparisons import ComparisonPair, comparison_pair
from redovisningai.accounting.periods import month, parse_period
from redovisningai.accounting.statements import StatementMapping
from redovisningai.analytics.differences import FAMILIES, DifferenceSet, collect_differences
from redovisningai.devdata.generator import DEMO_PROFILES, generate
from redovisningai.reports.comparison_report import ReportSelectionError, SelectedItem, build_comparison_report
from redovisningai.reports.document import to_docx, to_pdf, to_xlsx
from redovisningai.review.analysis import PAYROLL
from redovisningai.rules.rates import default_rates


@pytest.fixture(scope="module")
def bygg() -> LedgerIndex:
    return LedgerIndex.build(generate(DEMO_PROFILES[0], date(2026, 9, 30)).ledger)


def _differences(index: LedgerIndex, pair: ComparisonPair, **kw: object) -> DifferenceSet:
    return collect_differences(
        index,
        pair,
        mapping=StatementMapping(),
        categories=CategoryMapping(),
        rates=default_rates(),
        **kw,  # type: ignore[arg-type]
    )


def _report(index: LedgerIndex, pair: ComparisonPair, ds: DifferenceSet, items: list[SelectedItem], **kw: object):  # type: ignore[no-untyped-def]
    return build_comparison_report(
        index,
        pair,
        ds,
        items,
        company_name="Bygg & Co AB",
        org_number="556677-8899",
        mapping=StatementMapping(),
        rates=default_rates(),
        firm_name="Byrån",
        prepared_by="Anna",
        now=datetime(2026, 10, 1, 9, 0),
        **kw,  # type: ignore[arg-type]
    )


def _docx_text(data: bytes) -> str:
    d = Docx(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_differences_rank_and_recommend_diverse_items(bygg: LedgerIndex) -> None:
    pair = comparison_pair(parse_period("YTD:2026-09", bygg.ledger), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    kinds = {i.kind for i in ds.items}
    assert {"metric", "line", "category"} <= kinds
    recommended = ds.recommended()
    assert 1 <= len(recommended) <= 6
    families = [FAMILIES.get(i.id, i.id) for i in recommended]
    assert len(families) == len(set(families))
    assert all(i.selectable and i.score > 0 for i in recommended)
    scores = [i.score for i in ds.items if i.selectable]
    assert scores == sorted(scores, reverse=True)
    # Deterministiskt: samma underlag ger samma rangordning och texter.
    again = _differences(bygg, pair)
    assert [(i.id, i.score, i.summary) for i in again.items] == [(i.id, i.score, i.summary) for i in ds.items]
    consultants = ds.get("category:external_services")
    assert consultants is not None and consultants.previous == 0 and "ny post" in consultants.reason


def test_differences_hide_payroll_details_and_internal_findings(bygg: LedgerIndex) -> None:
    pair = comparison_pair(month(2026, 9), "yoy", bygg.ledger, bygg)
    hidden = _differences(bygg, pair, hidden_accounts=PAYROLL, include_findings=False)
    assert all(i.kind != "finding" for i in hidden.items)
    for item in hidden.items:
        for account in range(7000, 7700):
            assert f"konto {account}" not in item.summary
    personnel = hidden.get("line:income:personnel")
    assert personnel is not None and personnel.change is not None


def test_metric_summary_uses_computed_values(bygg: LedgerIndex) -> None:
    pair = comparison_pair(month(2026, 9), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    net_sales = ds.get("metric:net_sales")
    assert net_sales is not None
    assert net_sales.change == net_sales.current - net_sales.previous  # type: ignore[operator]
    assert net_sales.summary.startswith("Nettoomsättning ökade från 1,23")
    margin = ds.get("metric:operating_margin")
    assert margin is not None and "procentenheter" in margin.summary


def test_internal_report_contains_selection_comments_and_evidence(bygg: LedgerIndex) -> None:
    pair = comparison_pair(parse_period("YTD:2026-09", bygg.ledger), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    selection = [
        SelectedItem("metric:operating_margin", "Marginalen har förbättrats."),
        SelectedItem("line:income:other_external", "Konsultkostnaderna diskuteras på mötet."),
        SelectedItem("structure:operating_margin", None, "same_month", 3),
    ]
    doc, sheets = _report(bygg, pair, ds, selection, audience="internal")
    headings = [s.heading for s in doc.sections]
    assert headings[:3] == ["Om rapporten", "Sammanfattning", "Valda nyckeltal"]
    assert any(h.startswith("1. Rörelsemarginal") for h in headings)
    assert any(h.startswith("Största verifikationer") for h in headings)
    assert headings[-1] == "Bilaga: samtliga nyckeltal"
    assert doc.classification.startswith("INTERN")
    text = _docx_text(to_docx(doc))
    assert "Marginalen har förbättrats." in text and "Konsultkostnaderna diskuteras på mötet." in text
    assert "sep 2024" in text and "sep 2026" in text  # utveckling över tid
    assert to_pdf(doc).startswith(b"%PDF")
    wb = load_workbook(io.BytesIO(to_xlsx(sheets)))
    assert wb.sheetnames[0] == "Sammanfattning"
    assert len(wb.sheetnames) == 4


def test_metric_bridge_table_sums_to_change(bygg: LedgerIndex) -> None:
    pair = comparison_pair(month(2026, 9), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    _doc, sheets = _report(bygg, pair, ds, [SelectedItem("metric:operating_result")], audience="internal")
    table = dict(sheets)["1 Rörelseresultat"]
    effects = [row[3] for row in table.rows[:-1]]
    assert table.rows[-1][0] == "Summa förändring"
    assert sum(effects, Decimal(0)) == table.rows[-1][3]


def test_account_table_reconciles_with_line_total(bygg: LedgerIndex) -> None:
    pair = comparison_pair(parse_period("YTD:2026-09", bygg.ledger), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    for item in ds.items:
        if item.kind not in ("line", "category") or not item.selectable:
            continue
        _doc, sheets = _report(bygg, pair, ds, [SelectedItem(item.id)], audience="internal", include_key_figures=False)
        table = sheets[1][1]
        body, total = table.rows[:-1], table.rows[-1]
        assert total[1] == "Summa"
        assert sum((row[2] for row in body), Decimal(0)) == total[2], item.id
        assert sum((row[3] for row in body), Decimal(0)) == total[3], item.id


def test_client_report_excludes_findings_vouchers_and_payroll_accounts(bygg: LedgerIndex) -> None:
    pair = comparison_pair(month(2026, 9), "yoy", bygg.ledger, bygg)
    internal = _differences(bygg, pair)
    finding = next((i for i in internal.items if i.kind == "finding"), None)
    if finding is not None:
        with pytest.raises(ReportSelectionError, match="kundrapport"):
            _report(bygg, pair, internal, [SelectedItem(finding.id)], audience="client")
    client = _differences(bygg, pair, hidden_accounts=PAYROLL, include_findings=False)
    doc, sheets = _report(
        bygg, pair, client, [SelectedItem("line:income:personnel"), SelectedItem("metric:net_sales")], audience="client"
    )
    assert doc.classification == "KUNDRAPPORT"
    assert not any(s.heading.startswith("Största verifikationer") for s in doc.sections)
    assert not any("Beräkningsversion" in str(row) for s in doc.sections if s.table for row in s.table.rows)
    personnel = sheets[1][1]
    assert all(not (row[0] and 7000 <= int(row[0]) <= 7699) for row in personnel.rows)
    assert any("Lönekonton" in str(row[1]) for row in personnel.rows)


def test_selection_is_validated_against_server_items(bygg: LedgerIndex) -> None:
    pair = comparison_pair(month(2026, 9), "yoy", bygg.ledger, bygg)
    ds = _differences(bygg, pair)
    with pytest.raises(ReportSelectionError, match="Okänd"):
        _report(bygg, pair, ds, [SelectedItem("metric:hittepa")], audience="internal")
    with pytest.raises(ReportSelectionError, match="minst en"):
        _report(bygg, pair, ds, [], audience="internal")
    with pytest.raises(ReportSelectionError, match="för lång"):
        _report(bygg, pair, ds, [SelectedItem("metric:net_sales", "x" * 2001)], audience="internal")
    with pytest.raises(ReportSelectionError, match="Antal perioder"):
        _report(bygg, pair, ds, [SelectedItem("structure:net_sales", None, "months", 99)], audience="internal")
    missing_pair = comparison_pair(month(2019, 5), "yoy", bygg.ledger, bygg)
    missing = _differences(bygg, missing_pair)
    with pytest.raises(ReportSelectionError, match="saknar underlag"):
        _report(bygg, missing_pair, missing, [SelectedItem("metric:net_sales")], audience="internal")
