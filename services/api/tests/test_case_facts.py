"""Fakta som fynd och ärenden hänvisar till måste gå att verifiera och rendera.

Påhittad minimal SIE-fil: lokalhyran (5010) bokförs varje månad januari–augusti 2026 men
saknas i september. Kontrollen COST_DEVIATION ger då fyndet "Förväntad kostnad saknas" med
ett faktum för normalnivån, och ärendets trolig orsak hänvisar till det med {f:id}.
"""

from __future__ import annotations

import logging
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from redovisningai.ai.tasks import TASKS
from redovisningai.cli import main
from redovisningai.review.analysis import CompanyAnalysis, CompanyContext, ReviewResult
from redovisningai.rules.engine import CompanySettings
from redovisningai.sie.convert import ledger_from_documents
from redovisningai.sie.parser import parse_sie
from sie_samples import minimal_sie_with_missing_rent

PERIOD = "2026-09"


def _analysis() -> CompanyAnalysis:
    ledger = ledger_from_documents([(parse_sie(minimal_sie_with_missing_rent()), "minimal.se")])
    return CompanyAnalysis(ledger, CompanyContext("org", "company", ledger.company_name, CompanySettings()))


def _review(analysis: CompanyAnalysis) -> ReviewResult:
    return analysis.review(analysis.period(PERIOD))


def _rent_case(analysis: CompanyAnalysis, review: ReviewResult) -> dict:  # type: ignore[type-arg]
    return next(c for c in analysis.case_package(review)["cases"] if "5010" in c["title"])


def test_review_store_holds_the_facts_its_findings_cite() -> None:
    review = _review(_analysis())
    cited = {fact["id"] for record in review.records for fact in record.to_dict()["facts"]}
    assert cited, "den saknade hyran ska ge ett fynd med ett faktum"
    assert {fid for fid in cited if fid not in review.store} == set()


def test_review_store_includes_facts_saved_with_earlier_open_findings() -> None:
    # Ett öppet fynd från augusti ingår i septembers ärenden men dess kontroll körs inte om.
    analysis = _analysis()
    record = next(r for r in _review(analysis).records if r.facts)
    saved = {**record.facts[0], "id": "median_5010_from_august"}
    earlier = replace(record, fingerprint="found-in-august", period="2026-08", facts=[saved])

    review = analysis.review(analysis.period(PERIOD), existing=[earlier])

    assert "median_5010_from_august" in review.store


def test_unreadable_saved_fact_is_logged_and_skipped(caplog: pytest.LogCaptureFixture) -> None:
    analysis = _analysis()
    record = next(r for r in _review(analysis).records if r.facts)
    broken = {"id": "broken_fact", "unit": "no-such-unit"}
    earlier = replace(record, fingerprint="found-in-august", period="2026-08", facts=[broken])

    with caplog.at_level(logging.WARNING, logger="redovisningai.review.analysis"):
        review = analysis.review(analysis.period(PERIOD), existing=[earlier])

    assert "broken_fact" not in review.store
    assert "broken_fact" in caplog.text


def test_case_builder_accepts_claims_that_cite_the_findings_own_facts() -> None:
    analysis = _analysis()
    review = _review(analysis)
    case = _rent_case(analysis, review)
    fact_id = case["findings"][0]["facts"][0]["id"]
    answer = {
        "cases": [
            {
                "finding_ids": [f["id"] for f in case["findings"]],
                "title": case["title"],
                "root_cause": [
                    {
                        "type": "OBSERVATION",
                        "text": f"Hyran brukar vara {{f:{fact_id}}} per månad men saknas i perioden.",
                        "fact_ids": [fact_id],
                    }
                ],
                "suggestion": "ASK_CLIENT",
                "rationale": [
                    {"type": "OBSERVATION", "text": f"Normalnivån är {{f:{fact_id}}}.", "fact_ids": [fact_id]}
                ],
            }
        ]
    }
    package = analysis.case_package(review)
    _, result = TASKS["A2"].verify(answer, package, review.store, analysis.allowed_identifiers())
    assert [r.reason for r in result.rejected] == []


def test_analyze_excel_shows_case_causes_with_amounts_not_placeholders(tmp_path: Path) -> None:
    analysis = _analysis()
    normal_level = _rent_case(analysis, _review(analysis))["findings"][0]["facts"][0]["display"]
    sie = tmp_path / "minimal.se"
    sie.write_bytes(minimal_sie_with_missing_rent())
    out = tmp_path / "rapport"

    assert main(["analyze", str(sie), "--period", PERIOD, "--out", str(out)]) == 0

    sheet = load_workbook(next(out.glob("*granskning.xlsx")), read_only=True)["Ärenden"]
    causes = [str(row[2]) for row in sheet.iter_rows(min_row=2, values_only=True)]
    rent_cause = next(c for c in causes if "per månad" in c)
    assert "{f:" not in rent_cause
    assert normal_level in rent_cause
