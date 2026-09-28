"""Rapporternas läsordning och exporter med syntetisk bokföring."""

import re
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO

from docx import Document as WordDocument

from redovisningai.accounting.periods import month
from redovisningai.accounting.variance import BridgeComponent
from redovisningai.api.routes_other import _report_review_items
from redovisningai.devdata.generator import DEMO_PROFILES, generate
from redovisningai.reports.analysis import _result_summary, build_report_analysis
from redovisningai.reports.builders import client_report, internal_report
from redovisningai.reports.document import to_docx, to_pdf
from redovisningai.review.analysis import CompanyAnalysis, CompanyContext


def _synthetic_inputs():  # type: ignore[no-untyped-def]
    generated = generate(DEMO_PROFILES[0], date(2026, 9, 30))
    analysis = CompanyAnalysis(generated.ledger, CompanyContext("synthetic", "test", "Testbolag AB"))
    current, previous = month(2026, 9), month(2025, 9)
    return (
        analysis.overview(current, previous),
        analysis.statements(current, previous),
        build_report_analysis(analysis, current, previous),
    )


def test_customer_report_reads_as_analysis_then_actions_then_appendices() -> None:
    overview, statements, analysis = _synthetic_inputs()
    meeting = {
        "summary": [{"type": "EXPLANATION", "rendered": "Omsättningen steg genom fler bokförda fakturor."}],
        "questions": [{"type": "QUESTION", "rendered": "Ska ökningen följas upp i oktober?"}],
        "case_questions": [{"question": "Intern fråga om låneförbud"}],
    }
    report = client_report(overview, statements, meeting, "Testbyrån", analysis=analysis)
    headings = [section.heading for section in report.sections]
    assert headings == [
        "Avstämning av underlag",
        "Perioden i korthet",
        "Vad som ligger bakom förändringen",
        "Konsultens bedömning",
        "Frågor och nästa steg",
        "Periodens huvudtal",
        "Bilaga 1 · Nyckeltal",
        "Bilaga 2 · Resultaträkning",
        "Bilaga 3 · Balansräkning",
        "Om rapporten",
    ]
    assert "skiljer sig med 1 kr" in report.sections[0].paragraphs[0]
    assert "440" in report.sections[1].paragraphs[0]
    assert "Övriga resultatrader bidrog netto" in report.sections[1].paragraphs[0]
    assert "+896" in report.sections[2].paragraphs[0]
    assert "−518" in report.sections[2].paragraphs[1]
    assert "+63" in report.sections[2].paragraphs[3]
    assert report.sections[8].page_break_before  # balansräkningen börjar på ny sida
    assert all("låneförbud" not in str(section) for section in report.sections)
    assert report.sections[7].table is not None
    assert report.sections[7].table.headers[-1] == "Förändring, kr"
    assert all(
        "tkr" not in str(cell) and "Mkr" not in str(cell) for row in report.sections[7].table.rows for cell in row
    )

    word = WordDocument(BytesIO(to_docx(report)))
    word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "största delposten på intäktsraden" in word_text and "Omsättningen steg" in word_text
    assert "Bilaga 3 · Balansräkning" in word_text
    assert "låneförbud" not in word_text.lower()

    pdf = to_pdf(report)
    assert pdf.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page\b", pdf)) >= 2


def test_internal_report_puts_follow_up_before_traceable_findings() -> None:
    overview, statements, analysis = _synthetic_inputs()
    report = internal_report(
        overview,
        [
            {
                "rule_code": "TEST",
                "severity": "HIGH",
                "title": "Kontrollera bokning",
                "status": "NEW",
                "vouchers": ["A1"],
            }
        ],
        [{"severity": "HIGH", "title": "Öppet ärende", "status": "OPEN", "suggested_action": "Stäm av fakturan."}],
        {
            "claims": [
                {"type": "HYPOTHESIS", "rendered": "Möjlig förklaring: En stor bokning kan förklara förändringen."}
            ]
        },
        {"status": "COMPLETE", "accounting_method": "invoice", "notes": []},
        statements=statements,
        analysis=analysis,
    )
    headings = [section.heading for section in report.sections]
    assert headings[:6] == [
        "Avstämning av underlag",
        "Perioden i korthet",
        "Vad som ligger bakom förändringen",
        "AI-stödd kommentar för intern granskning",
        "Prioriterad uppföljning",
        "Underlag och periodmognad",
    ]
    assert headings[-2:] == ["Bilaga 4 · Resultaträkning", "Bilaga 5 · Balansräkning"]
    assert report.sections[10].page_break_before  # balansräkningen börjar på ny sida
    assert "Möjlig förklaring:" in report.sections[3].paragraphs[0]
    assert report.sections[3].paragraphs[0].count("Möjlig förklaring:") == 1
    assert "A1" in str(report.sections[8].table.rows)


def test_internal_report_excludes_restricted_aml_cases_and_findings() -> None:
    overview, statements, analysis = _synthetic_inputs()
    report = internal_report(
        overview,
        [{"visibility": "RESTRICTED_AML", "title": "PTL-hemligt"}],
        [{"visibility": "RESTRICTED_AML", "title": "PTL-hemligt", "status": "OPEN"}],
        None,
        {"status": "COMPLETE", "accounting_method": "invoice", "notes": []},
        statements=statements,
        analysis=analysis,
    )
    assert "PTL-hemligt" not in str(report)
    assert "Inga öppna ärenden" in str(report)


def test_balanced_period_has_no_reconciliation_warning() -> None:
    overview, original, analysis = _synthetic_inputs()
    statements = deepcopy(original)
    totals = {line["code"]: line for line in statements["balance"]["lines"]}
    totals["total_equity_liabilities"]["amount"] = totals["total_assets"]["amount"]
    report = client_report(overview, statements, None, "Testbyrån", analysis=analysis)
    assert report.sections[0].heading == "Perioden i korthet"
    assert "Rörelseresultatet" in report.sections[0].paragraphs[0]


def test_worsening_period_describes_negative_drivers_as_causes() -> None:
    summary = _result_summary(
        Decimal("80"),
        Decimal("100"),
        [
            BridgeComponent("net_sales", "Nettoomsättning", Decimal(0), Decimal(0), Decimal("-30")),
            BridgeComponent("materials", "Råvaror", Decimal(0), Decimal(0), Decimal("10")),
        ],
    )
    assert "försämrades" in summary
    assert "största negativa bokföringsbidragen" in summary
    assert "dämpade försämringen" in summary


def test_summary_names_every_selected_driver_before_calculating_other_rows() -> None:
    components = [
        BridgeComponent(code, label, Decimal(0), Decimal(0), Decimal(10))
        for code, label in (
            ("net_sales", "Nettoomsättning"),
            ("materials", "Råvaror"),
            ("personnel", "Personalkostnader"),
            ("other_operating_income", "Övriga rörelseintäkter"),
        )
    ]
    summary = _result_summary(Decimal(40), Decimal(0), components)
    assert "övriga rörelseintäkter" in summary
    assert "Övriga resultatrader" not in summary


def test_unreviewed_cli_customer_report_is_marked_as_draft_inside_document() -> None:
    overview, statements, analysis = _synthetic_inputs()
    meeting = {"summary": [{"type": "EXPLANATION", "rendered": "Omsättningen har förändrats."}]}
    report = client_report(overview, statements, meeting, "Testbyrån", analysis=analysis, draft=True)
    assert report.classification == "UTKAST – FÅR INTE LÄMNAS TILL KUND"
    assert "AI-utkast för konsultens granskning" in [section.heading for section in report.sections]
    assert "inte godkänd" in str(report.sections)


def test_internal_report_items_follow_payroll_permission() -> None:
    payroll = {
        "rule_code": "EMPLOYER_CONTRIBUTION_RATIO",
        "accounts": [7510],
        "title": "Löneavvikelse",
        "vouchers": ["L1"],
    }
    ordinary = {"rule_code": "ACCOUNT_CHANGE", "accounts": [6550], "title": "Konsultarvode", "vouchers": ["A1"]}
    cases = [
        {"title": "Löneärende", "findings": [payroll]},
        {"title": "Kostnadsärende", "findings": [ordinary]},
    ]
    findings, visible_cases = _report_review_items([payroll, ordinary], cases, can_payroll=False)
    assert findings == [ordinary]
    assert visible_cases == [cases[1]]
    assert _report_review_items([payroll, ordinary], cases, can_payroll=True) == ([payroll, ordinary], cases)
