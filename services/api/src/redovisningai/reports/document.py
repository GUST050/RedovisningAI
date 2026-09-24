"""Gemensam dokumentmodell och rendering till Word (docx), PDF och Excel (xlsx)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any


@dataclass(slots=True)
class Table:
    headers: list[str]
    rows: list[list[Any]]
    numeric_cols: set[int] = field(default_factory=set)


@dataclass(slots=True)
class Section:
    heading: str
    paragraphs: list[str] = field(default_factory=list)
    bullets: list[str] = field(default_factory=list)
    table: Table | None = None
    note: str | None = None


@dataclass(slots=True)
class Document:
    title: str
    subtitle: str
    sections: list[Section]
    footer: str = ""
    classification: str = ""  # t.ex. "INTERN" / "KUND"
    created_at: datetime = field(default_factory=datetime.now)


# ---------------------------------------------------------------------------- Word


def to_docx(doc: Document) -> bytes:
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    d = Docx()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    if doc.classification:
        p = d.add_paragraph(doc.classification)
        p.runs[0].font.size = Pt(8)
        p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    d.add_heading(doc.title, level=0)
    sub = d.add_paragraph(doc.subtitle)
    sub.runs[0].italic = True
    for s in doc.sections:
        d.add_heading(s.heading, level=1)
        for para in s.paragraphs:
            d.add_paragraph(para)
        for b in s.bullets:
            d.add_paragraph(b, style="List Bullet")
        if s.table is not None and s.table.rows:
            t = d.add_table(rows=1, cols=len(s.table.headers))
            t.style = "Light Grid Accent 1"
            for i, h in enumerate(s.table.headers):
                t.rows[0].cells[i].text = str(h)
            for row in s.table.rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = "" if v is None else str(v)
                    if i in s.table.numeric_cols:
                        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if s.note:
            n = d.add_paragraph(s.note)
            n.runs[0].font.size = Pt(8.5)
            n.runs[0].italic = True
    if doc.footer:
        f = d.add_paragraph(doc.footer)
        f.runs[0].font.size = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------- PDF


def to_pdf(doc: Document) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from reportlab.platypus import Table as RLTable

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=doc.title,
    )
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=colors.grey)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=10.5)
    cell_r = ParagraphStyle("cellr", parent=cell, alignment=2)

    def esc(t: Any) -> str:
        return (str(t) if t is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list[Any] = []
    if doc.classification:
        story.append(Paragraph(esc(doc.classification), small))
    story += [Paragraph(esc(doc.title), ss["Title"]), Paragraph(esc(doc.subtitle), small), Spacer(1, 6)]
    for s in doc.sections:
        story.append(Paragraph(esc(s.heading), ss["Heading2"]))
        for para in s.paragraphs:
            story.append(Paragraph(esc(para), body))
        if s.bullets:
            story.append(
                ListFlowable([ListItem(Paragraph(esc(b), body)) for b in s.bullets], bulletType="bullet", leftIndent=12)
            )
        if s.table is not None and s.table.rows:
            data = [[Paragraph(f"<b>{esc(h)}</b>", cell) for h in s.table.headers]]
            for row in s.table.rows:
                data.append(
                    [Paragraph(esc(v), cell_r if i in s.table.numeric_cols else cell) for i, v in enumerate(row)]
                )
            t = RLTable(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C5CDD8")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [Spacer(1, 4), t]
        if s.note:
            story.append(Paragraph(esc(s.note), small))
        story.append(Spacer(1, 6))
    if doc.footer:
        story.append(Paragraph(esc(doc.footer), small))
    pdf.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------- Excel


def to_xlsx(sheets: list[tuple[str, Table]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for name, table in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(table.headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8EDF3")
        for row in table.rows:
            ws.append([_excel_value(v) for v in row])
        for idx in table.numeric_cols:
            for r in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
                for c in r:
                    c.number_format = "#,##0.00"
                    c.alignment = Alignment(horizontal="right")
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _excel_value(v: Any) -> Any:
    from decimal import Decimal, InvalidOperation

    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, str):
        try:
            if v and v.lstrip("-").replace(".", "", 1).isdigit():
                return float(Decimal(v))
        except InvalidOperation:
            pass
        # Skydd mot formelinjektion i Excel (CSV/XLSX-injection).
        if v[:1] in ("=", "+", "-", "@") and not v[1:2].isdigit():
            return "'" + v
    return v
