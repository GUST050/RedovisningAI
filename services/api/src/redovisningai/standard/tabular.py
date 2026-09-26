"""Verifikationslistor och huvudböcker i CSV- eller Excelformat (.xlsx).

Många bokföringsprogram, bland annat Fortnox, kan exportera verifikationer som tabell. Kolumnerna
heter olika i olika program, så de känns igen på rubriken ("Vernr", "Verifikationsnummer",
"Konto", "Debet", "Kredit", "Belopp" …). Både platta listor (en rad per konteringsrad) och
listor där verifikationsnummer eller konto bara står på första raden i en grupp stöds.

Principer (samma som för SIE):
- Filen avvisas bara om den inte går att tolka alls. Allt annat blir `ParseIssue` med radnummer.
- Belopp blir Decimal. Svenska format ("1 234,50", "−500", "(500)", "500-") förstås.
- Det som inte finns i filen hittas inte på: saknas ingående balanser markeras det, och bara de
  månader som filen täcker räknas som kända.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from datetime import date as Date  # alias: fältet _VoucherDraft.date skuggar typen i klassen
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from itertools import pairwise

from redovisningai.accounting.periods import add_months, month_end, month_start, months_between
from redovisningai.domain.ledger import ZERO, Account, FiscalYear, Ledger, Row, Voucher, YearData
from redovisningai.sie.parser import ParseIssue, Severity

TABULAR_PARSER_VERSION = "1.0.0"
MAX_ROWS = 2_000_000
MAX_COLUMNS = 200
HEADER_SCAN_ROWS = 40
MAX_ROW_MESSAGES = 50


class TabularFormatError(ValueError):
    """Filen går inte att tolka som en verifikationslista."""


# ---------------------------------------------------------------------------- kolumnigenkänning

FIELD_LABELS = {
    "voucher": "Verifikation",
    "series": "Serie",
    "number": "Verifikationsnummer",
    "date": "Datum",
    "account": "Konto",
    "account_name": "Kontonamn",
    "debit": "Debet",
    "credit": "Kredit",
    "amount": "Belopp",
    "text": "Text",
    "row_text": "Radtext",
    "cost_center": "Kostnadsställe",
    "project": "Projekt",
    "quantity": "Antal",
    "registered": "Registreringsdatum",
    "signature": "Signatur",
    "balance": "Saldo",
}

_SYNONYMS: dict[str, tuple[str, ...]] = {
    "voucher": (
        "vernr",
        "verifikation",
        "verifikationsnummer",
        "verifikationsnr",
        "vernummer",
        "verifikat",
        "verifikationsid",
        "ver",
        "voucher",
        "vouchernumber",
        "voucherno",
        "vouchernr",
        "voucherid",
    ),
    "series": ("serie", "verifikationsserie", "verserie", "series", "voucherseries"),
    "number": ("nummer", "nr", "lopnr", "lopnummer", "number"),
    "date": (
        "datum",
        "bokforingsdatum",
        "bokfdatum",
        "verifikationsdatum",
        "verdatum",
        "vdatum",
        "transaktionsdatum",
        "transdatum",
        "date",
        "voucherdate",
        "transactiondate",
        "bookingdate",
        "postingdate",
        "accountingdate",
    ),
    "account": ("konto", "kontonr", "kontonummer", "kto", "account", "accountnumber", "accountno", "accountnr"),
    "account_name": (
        "kontonamn",
        "kontobenamning",
        "benamning",
        "kontobeskrivning",
        "accountname",
        "accountdescription",
    ),
    "debit": ("debet", "debit", "debetbelopp", "debetsek", "debetkr"),
    "credit": ("kredit", "credit", "kreditbelopp", "kreditsek", "kreditkr"),
    "amount": ("belopp", "amount", "summa", "beloppsek", "beloppkr", "transaktionsbelopp", "value"),
    "text": (
        "text",
        "beskrivning",
        "verifikationstext",
        "vertext",
        "verifikationsbeskrivning",
        "description",
        "vouchertext",
        "voucherdescription",
        "rubrik",
    ),
    "row_text": (
        "transaktionsinfo",
        "transinfo",
        "transaktionsinformation",
        "transaktionstext",
        "radtext",
        "radbeskrivning",
        "specifikation",
        "transactioninformation",
        "transactiontext",
    ),
    "cost_center": ("kostnadsstalle", "ks", "kst", "kostnadsst", "costcenter", "resultatenhet", "resultatstalle"),
    "project": ("projekt", "projektnr", "projnr", "proj", "projektnummer", "project"),
    "quantity": ("antal", "kvantitet", "quantity", "qty"),
    "registered": ("registreringsdatum", "regdatum", "registrerad", "skapad", "createddate", "registrationdate"),
    "signature": ("signatur", "anvandare", "registreradav", "skapadav", "user", "createdby", "signature"),
    "balance": ("saldo", "balance", "ackumuleratsaldo", "utgaendesaldo"),
}
_LOOKUP = {syn: fld for fld, syns in _SYNONYMS.items() for syn in syns}
_TRANS = str.maketrans({"å": "a", "ä": "a", "ö": "o", "é": "e", "è": "e", "ü": "u"})


def normalize_header(text: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text).lower().translate(_TRANS))


@dataclass(frozen=True, slots=True)
class ColumnMap:
    fields: dict[str, int]
    header_row: int  # 1-baserat radnummer i filen/bladet
    headers: tuple[str, ...]
    sheet: str | None = None

    def describe(self) -> str:
        parts = [f"{FIELD_LABELS[f]} ← ”{self.headers[i]}”" for f, i in sorted(self.fields.items(), key=lambda x: x[1])]
        return ", ".join(parts)


def _map_header(cells: Sequence[object]) -> tuple[dict[str, int], list[str]]:
    fields: dict[str, int] = {}
    duplicates: list[str] = []
    for i, cell in enumerate(cells[:MAX_COLUMNS]):
        if cell is None:
            continue
        fld = _LOOKUP.get(normalize_header(cell))
        if fld is None:
            continue
        if fld in fields:
            duplicates.append(str(cell))
            continue
        fields[fld] = i
    return fields, duplicates


def _is_usable(fields: dict[str, int]) -> bool:
    has_amount = any(f in fields for f in ("amount", "debit", "credit"))
    has_voucher = "voucher" in fields or "number" in fields
    return "account" in fields and has_amount and has_voucher and "date" in fields


def find_header(rows: Sequence[Sequence[object]]) -> tuple[int, dict[str, int], list[str]] | None:
    """Hitta rubrikraden bland de första raderna. Returnerar (index, fält, dubbletter)."""
    best: tuple[int, dict[str, int], list[str]] | None = None
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        fields, dups = _map_header(row)
        if _is_usable(fields) and (best is None or len(fields) > len(best[1])):
            best = (i, fields, dups)
    return best


# ---------------------------------------------------------------------------- värden


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


_SPACES = re.compile(r"[\s   ']")
_CURRENCY = re.compile(r"(?i)(sek|kr\.?|:-)$")


def parse_amount(value: object) -> Decimal | None:
    """Tolka ett belopp. Tom cell ger None. Kastar ValueError för text som inte är ett belopp."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("sant/falskt är inget belopp")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        exact = Decimal(repr(value))
        rounded = exact.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        return rounded if abs(exact - rounded) < Decimal("0.000001") else exact
    s = _SPACES.sub("", str(value))
    if not s or s in ("-", "–", "—"):
        return None
    s = _CURRENCY.sub("", s)
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1]
    if s.endswith("-") and len(s) > 1:
        negative, s = True, s[:-1]
    s = s.replace("−", "-").replace("–", "-")
    if s.startswith("+"):
        s = s[1:]
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    elif s.count(",") > 1 or s.count(".") > 1:
        s = s.replace(",", "").replace(".", "")
    if not re.fullmatch(r"-?\d+(\.\d+)?", s):
        raise ValueError(f"ogiltigt belopp {value!r}")
    try:
        amount = Decimal(s)
    except InvalidOperation as exc:  # pragma: no cover – skyddas av mönstret ovan
        raise ValueError(f"ogiltigt belopp {value!r}") from exc
    return -amount if negative else amount


_ISO = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:[ T].*)?$")
_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
_EUROPEAN = re.compile(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$")
_EXCEL_EPOCH = date(1899, 12, 30)


def parse_day(value: object) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        n = int(value)
        if 19000101 <= n <= 21001231:
            return parse_day(str(n))
        if 20000 <= n <= 80000:
            return _EXCEL_EPOCH + timedelta(days=n)
        raise ValueError(f"ogiltigt datum {value!r}")
    s = str(value).strip()
    if not s:
        return None
    for pattern, order in ((_ISO, (1, 2, 3)), (_COMPACT, (1, 2, 3)), (_EUROPEAN, (3, 2, 1))):
        m = pattern.match(s)
        if m:
            y, mo, d = (int(m.group(i)) for i in order)
            try:
                return date(y, mo, d)
            except ValueError as exc:
                raise ValueError(f"ogiltigt datum {value!r}") from exc
    raise ValueError(f"ogiltigt datum {value!r}")


_ACCOUNT = re.compile(r"^\s*(\d{3,8})(?:\s*[-–:]?\s+(.+))?\s*$")


def parse_account(value: object) -> tuple[int, str | None] | None:
    if value is None:
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        return (value, None) if value > 0 else None
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value), None
    s = str(value).strip()
    if not s:
        return None
    m = _ACCOUNT.match(s)
    if not m:
        raise ValueError(f"ogiltigt konto {value!r}")
    return int(m.group(1)), (m.group(2) or None)


_VOUCHER = re.compile(r"^([A-Za-zÅÄÖåäö]{1,10})\s*[-:.]?\s*(\d+)$")


def split_voucher(text: str) -> tuple[str, str]:
    s = text.strip()
    m = _VOUCHER.match(s)
    if m:
        return m.group(1).upper(), m.group(2)
    if s.isdigit():
        return "", s
    return "", s


_IB = re.compile(r"^(ib|ing[aå]ende\s*(balans|saldo)|opening\s*balance)\b", re.I)
_UB = re.compile(r"^(ub|utg[aå]ende\s*(balans|saldo)|closing\s*balance)\b", re.I)
_SUMMARY = re.compile(r"^(summa|total|totalt|omslutning|delsumma|periodens\s*(förändring|summa))\b", re.I)
_ORGNR = re.compile(r"(?<!\d)(\d{6}-\d{4})(?!\d)")
_ORG_LABEL = re.compile(r"(?i)(org(anisations)?\.?\s*(nr|nummer)\.?:?)\s*$")
_RANGE = re.compile(
    r"(\d{4}-\d{2}-\d{2})\s*(?:-|–|—|till|to|t\.?o\.?m\.?)\s*(\d{4}-\d{2}-\d{2})",
    re.I,
)


# ---------------------------------------------------------------------------- läsning av filer


def decode_text(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace"), "utf-16"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace"), "cp1252"


def _sniff_delimiter(lines: Sequence[str]) -> str:
    best = (-1, -1, ";")
    for cand in (";", "\t", ",", "|"):
        counts = [len(r) for r in csv.reader(lines, delimiter=cand) if len(r) > 1]
        if not counts:
            continue
        mode, freq = Counter(counts).most_common(1)[0]
        score = (freq, mode, cand)
        if score[:2] > best[:2]:
            best = score
    return best[2]


def read_csv_rows(raw: bytes) -> tuple[list[list[str]], str, str]:
    text, encoding = decode_text(raw)
    lines = text.splitlines()
    sample = [ln for ln in lines[:200] if ln.strip()][:60]
    delimiter = _sniff_delimiter(sample)
    rows: list[list[str]] = []
    for row in csv.reader(lines, delimiter=delimiter):
        rows.append(row[:MAX_COLUMNS])
        if len(rows) > MAX_ROWS:
            raise TabularFormatError(f"Filen har fler än {MAX_ROWS:,} rader – dela upp exporten per räkenskapsår.")
    return rows, encoding, delimiter


def read_xlsx_rows(raw: bytes) -> list[tuple[str, list[list[object]]]]:
    """Alla blad som (namn, rader). Formler läses som sparade värden."""
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise TabularFormatError(f"Excelfilen går inte att läsa ({exc.__class__.__name__}).") from exc
    sheets: list[tuple[str, list[list[object]]]] = []
    total = 0
    try:
        for ws in wb.worksheets:
            rows: list[list[object]] = []
            for values in ws.iter_rows(values_only=True, max_col=MAX_COLUMNS):
                rows.append(list(values))
                total += 1
                if total > MAX_ROWS:
                    raise TabularFormatError(f"Filen har fler än {MAX_ROWS:,} rader – dela upp exporten.")
            sheets.append((ws.title, rows))
    finally:
        wb.close()
    return sheets


# ---------------------------------------------------------------------------- tolkning


@dataclass(slots=True)
class TabularResult:
    ledger: Ledger
    issues: list[ParseIssue]
    columns: ColumnMap
    format: str  # CSV | XLSX
    encoding: str | None
    delimiter: str | None
    # Räkenskapsårets start → (första, sista) dag som filen täcker. Används för att en import av
    # en del av året bara ska ersätta verifikationer inom samma datumintervall.
    ranges: dict[date, tuple[date, date]] = field(default_factory=dict)


@dataclass(slots=True)
class _VoucherDraft:
    series: str
    number: str
    fiscal_year: FiscalYear
    line: int
    date: Date | None = None
    text: str = ""
    reg_date: Date | None = None
    signature: str | None = None
    rows: list[tuple[Row, Date]] = field(default_factory=list)


def fiscal_year_resolver(known: Sequence[FiscalYear], start_month: int | None = None) -> Callable[[date], FiscalYear]:
    """Räkenskapsår för ett datum: bolagets kända år först, annars år om tolv månader intill dem."""
    years = sorted(known, key=lambda fy: fy.start)

    def resolve(d: date) -> FiscalYear:
        for fy in years:
            if fy.contains(d):
                return fy
        if years:
            if d > years[-1].end:
                start = years[-1].end + timedelta(days=1)
                while True:
                    end = add_months(start, 12) - timedelta(days=1)
                    if d <= end:
                        return FiscalYear(start, end)
                    start = end + timedelta(days=1)
            if d < years[0].start:
                end = years[0].start - timedelta(days=1)
                while True:
                    start = add_months(end + timedelta(days=1), -12)
                    if d >= start:
                        return FiscalYear(start, end)
                    end = start - timedelta(days=1)
            for before, after in pairwise(years):
                if before.end < d < after.start:
                    return FiscalYear(before.end + timedelta(days=1), after.start - timedelta(days=1))
        sm = start_month or 1
        start = date(d.year if d.month >= sm else d.year - 1, sm, 1)
        return FiscalYear(start, add_months(start, 12) - timedelta(days=1))

    return resolve


def _preamble_info(rows: Sequence[Sequence[object]]) -> tuple[str | None, str | None, tuple[date, date] | None]:
    name: str | None = None
    org: str | None = None
    declared: tuple[date, date] | None = None
    for row in rows:
        for cell in row:
            text = _text(cell)
            if not text:
                continue
            if org is None and (m := _ORGNR.search(text)):
                org = m.group(1)
                before = _ORG_LABEL.sub("", text[: m.start()]).strip(" ,;:-–(")
                if before:
                    name = before
            if declared is None and (r := _RANGE.search(text)):
                try:
                    a, b = date.fromisoformat(r.group(1)), date.fromisoformat(r.group(2))
                except ValueError:
                    continue
                if a <= b:
                    declared = (a, b)
    return name, org, declared


def parse_rows(
    rows: Sequence[Sequence[object]],
    *,
    fmt: str,
    encoding: str | None = None,
    delimiter: str | None = None,
    sheet: str | None = None,
    fiscal_years: Sequence[FiscalYear] = (),
    fiscal_year_start_month: int | None = None,
) -> TabularResult:
    header = find_header(rows)
    if header is None:
        seen = [str(c) for r in rows[:HEADER_SCAN_ROWS] for c in r if c not in (None, "")][:12]
        raise TabularFormatError(
            "Hittar ingen rubrikrad med verifikation, datum, konto och belopp (eller debet/kredit). "
            + (f"Första värdena i filen: {', '.join(seen)}." if seen else "Filen verkar vara tom.")
        )
    hidx, fields, duplicates = header
    headers = tuple(_text(c) for c in rows[hidx])
    columns = ColumnMap(fields, hidx + 1, headers, sheet)
    issues: list[ParseIssue] = [
        ParseIssue(hidx + 1, "COLUMNS_DETECTED", f"Tolkade kolumner: {columns.describe()}.", Severity.INFO)
    ]
    for dup in duplicates:
        issues.append(
            ParseIssue(hidx + 1, "DUPLICATE_COLUMN", f"Kolumnen ”{dup}” används inte (dubblett).", Severity.INFO)
        )
    name, org, declared = _preamble_info(rows[:hidx])
    resolve = fiscal_year_resolver(fiscal_years, fiscal_year_start_month)

    def cell(row: Sequence[object], fld: str) -> object:
        i = fields.get(fld)
        if i is None or i >= len(row):
            return None
        v = row[i]
        return None if isinstance(v, str) and not v.strip() else v

    accounts: dict[int, str] = {}
    drafts: dict[tuple[date, str, str], _VoucherDraft] = {}
    openings: dict[date, dict[int, Decimal]] = {}
    objects: dict[tuple[str, str], str] = {}
    row_messages = 0
    skipped_summary = 0
    skipped_ub = 0
    ctx_voucher: tuple[str, str] | None = None
    ctx_date: date | None = None
    ctx_text = ""
    ctx_account: int | None = None

    def problem(line: int, message: str) -> None:
        nonlocal row_messages
        row_messages += 1
        if row_messages <= MAX_ROW_MESSAGES:
            issues.append(ParseIssue(line, "INVALID_ROW", message))

    for offset, row in enumerate(rows[hidx + 1 :], start=hidx + 2):
        if not any(_text(v) for v in row):
            continue
        try:
            account_raw = parse_account(cell(row, "account"))
        except ValueError:
            # Text i kontokolumnen (t.ex. "Summa konto 1930") är en summerings- eller rubrikrad.
            if _text(cell(row, "voucher")) or _text(cell(row, "number")):
                problem(offset, f"Raden hoppas över: ogiltigt konto {_text(cell(row, 'account'))!r}.")
            else:
                skipped_summary += 1
            continue
        try:
            day = parse_day(cell(row, "date"))
            amount = parse_amount(cell(row, "amount"))
            debit = parse_amount(cell(row, "debit"))
            credit = parse_amount(cell(row, "credit"))
            balance = parse_amount(cell(row, "balance"))
            quantity = parse_amount(cell(row, "quantity"))
            registered = parse_day(cell(row, "registered"))
        except ValueError as exc:
            problem(offset, f"Raden hoppas över: {exc}.")
            continue
        voucher_text = _text(cell(row, "voucher"))
        series_text = _text(cell(row, "series"))
        number_text = _text(cell(row, "number"))
        text = _text(cell(row, "text"))
        row_text = _text(cell(row, "row_text"))
        voucher_id: tuple[str, str] | None = None
        if voucher_text:
            series, number = split_voucher(voucher_text)
            voucher_id = (series or series_text.upper(), number)
        elif number_text:
            voucher_id = (series_text.upper(), number_text)
        # IB/UB känns igen på verifikationskolumnen ("IB") eller, för rader utan verifikation, på texten.
        marker = voucher_text or series_text
        is_ib = bool(_IB.match(marker)) or (voucher_id is None and bool(_IB.match(text)))
        is_ub = bool(_UB.match(marker)) or (voucher_id is None and bool(_UB.match(text)))
        if is_ib or is_ub:
            voucher_id = None
        has_amount = amount is not None or debit is not None or credit is not None
        if amount is None and (debit is not None or credit is not None):
            amount = (debit or ZERO) - (credit or ZERO)

        if account_raw is not None:
            account, inline_name = account_raw
            name_cell = _text(cell(row, "account_name")) or inline_name
            if name_cell or account not in accounts:
                accounts[account] = name_cell or accounts.get(account, "")
        else:
            account = None

        # Rad med konto men utan belopp och verifikation = kontorubrik i en huvudbok.
        if account is not None and not has_amount and voucher_id is None and not is_ib and balance is None:
            ctx_account = account
            continue
        # Rad med verifikation men utan konto och belopp = verifikationsrubrik.
        if account is None and not has_amount and voucher_id is not None and not is_ib:
            ctx_voucher, ctx_date, ctx_text = voucher_id, day, text
            continue

        if is_ub:
            skipped_ub += 1
            continue
        if is_ib:
            ib_account = account if account is not None else ctx_account
            ib_amount = amount if amount is not None else balance
            if ib_account is None or ib_amount is None:
                problem(offset, "IB-rad saknar konto eller belopp.")
                continue
            if ib_account >= 3000:
                problem(offset, f"IB-rad på resultatkonto {ib_account} ignoreras.")
                continue
            fy = resolve(day) if day else None
            key = fy.start if fy else date.min
            openings.setdefault(key, {})
            openings[key][ib_account] = openings[key].get(ib_account, ZERO) + ib_amount
            continue
        if not has_amount or amount is None:
            if account is not None and voucher_id is not None:
                problem(offset, "Raden saknar belopp.")
            continue
        if account is None:
            if voucher_id is None or ctx_account is None or _SUMMARY.match(text):
                skipped_summary += 1
                continue
            account = ctx_account
        if _SUMMARY.match(text) and voucher_id is None:
            skipped_summary += 1
            continue
        if voucher_id is None:
            if ctx_voucher is None:
                problem(offset, "Raden saknar verifikationsnummer.")
                continue
            voucher_id = ctx_voucher
        elif voucher_id != ctx_voucher:
            ctx_voucher, ctx_date, ctx_text = voucher_id, None, ""
        if day is None:
            day = ctx_date
        if day is None:
            problem(offset, f"Verifikation {''.join(voucher_id)} saknar datum.")
            continue
        if ctx_date is None:
            ctx_date = day
        fy = resolve(day)
        draft_key = (fy.start, voucher_id[0], voucher_id[1])
        draft = drafts.get(draft_key)
        if draft is None:
            draft = drafts[draft_key] = _VoucherDraft(voucher_id[0], voucher_id[1], fy, offset)
            draft.date = ctx_date if ctx_date and fy.contains(ctx_date) else day
            draft.text = ctx_text or text
            draft.reg_date = registered
            draft.signature = _text(cell(row, "signature")) or None
        elif not draft.text and (ctx_text or text):
            draft.text = ctx_text or text
        objs: list[tuple[str, str]] = []
        if cc := _text(cell(row, "cost_center")):
            objs.append(("1", cc))
            objects.setdefault(("1", cc), cc)
        if pr := _text(cell(row, "project")):
            objs.append(("6", pr))
            objects.setdefault(("6", pr), pr)
        row_label = row_text or (text if text and text != draft.text else "")
        draft.rows.append(
            (
                Row(
                    account=account,
                    amount=amount,
                    text=row_label or None,
                    quantity=quantity,
                    objects=tuple(objs),
                    source_line=offset,
                ),
                day,
            )
        )

    if row_messages > MAX_ROW_MESSAGES:
        issues.append(
            ParseIssue(None, "INVALID_ROW", f"Ytterligare {row_messages - MAX_ROW_MESSAGES} rader kunde inte tolkas.")
        )
    if skipped_summary:
        issues.append(
            ParseIssue(
                None,
                "SKIPPED_ROWS",
                f"{skipped_summary} summerings- eller textrader utan konto hoppades över.",
                Severity.INFO,
            )
        )
    if skipped_ub:
        issues.append(
            ParseIssue(
                None, "SKIPPED_UB", f"{skipped_ub} rader med utgående balans används inte (räknas fram).", Severity.INFO
            )
        )
    if not drafts and not openings:
        raise TabularFormatError("Filen innehåller inga konteringsrader med konto, datum och belopp.")

    ledger = Ledger(company_name=name or "Okänt bolag", org_number=org)
    if "cost_center" in fields:
        ledger.dimensions["1"] = "Kostnadsställe"
    if "project" in fields:
        ledger.dimensions["6"] = "Projekt"
    ledger.objects.update(objects)
    for acc_no, acc_name in accounts.items():
        ledger.accounts[acc_no] = Account(acc_no, acc_name or f"Konto {acc_no}")

    by_year: dict[date, list[Voucher]] = {}
    years: dict[date, FiscalYear] = {}
    unbalanced = 0
    for draft in drafts.values():
        assert draft.date is not None
        voucher = Voucher(
            series=draft.series,
            number=draft.number,
            date=draft.date,
            text=draft.text,
            rows=tuple(
                Row(
                    account=r.account,
                    amount=r.amount,
                    trans_date=None if day == draft.date else day,
                    text=r.text,
                    quantity=r.quantity,
                    objects=r.objects,
                    source_line=r.source_line,
                )
                for r, day in draft.rows
            ),
            reg_date=draft.reg_date,
            signature=draft.signature,
            source_line=draft.line,
        )
        if voucher.balance != 0:
            unbalanced += 1
            if unbalanced <= MAX_ROW_MESSAGES:
                issues.append(
                    ParseIssue(
                        draft.line,
                        "UNBALANCED_VOUCHER",
                        f"Verifikation {voucher.key} balanserar inte (differens {voucher.balance}).",
                    )
                )
        by_year.setdefault(draft.fiscal_year.start, []).append(voucher)
        years[draft.fiscal_year.start] = draft.fiscal_year
    if drafts and unbalanced > max(3, len(drafts) // 5):
        issues.append(
            ParseIssue(
                None,
                "MANY_UNBALANCED",
                f"{unbalanced} av {len(drafts)} verifikationer balanserar inte. Exporten verkar sakna rader "
                "(t.ex. filtrerad på konton eller period) – resultat- och balansräkning blir inte tillförlitliga.",
            )
        )
    undated_ib = openings.pop(date.min, None)
    if undated_ib:
        first = min(years) if years else None
        if first is None:
            raise TabularFormatError("Filen innehåller bara IB-rader utan datum och inga verifikationer.")
        target = openings.setdefault(first, {})
        for acc, amt in undated_ib.items():
            target[acc] = target.get(acc, ZERO) + amt
    for start in openings:
        if start not in years:
            years[start] = resolve(start)

    ranges: dict[date, tuple[date, date]] = {}
    for start in sorted(years):
        fy = years[start]
        vouchers = sorted(
            by_year.get(start, []),
            key=lambda v: (v.date, v.series, int(v.number) if v.number.isdigit() else 0, v.number),
        )
        dates = [v.date for v in vouchers]
        if declared is not None and dates and declared[0] <= min(dates) and max(dates) <= declared[1]:
            lo, hi = max(declared[0], fy.start), min(declared[1], fy.end)
        elif dates:
            # Utan angiven period gäller filen för hela de månader den har verifikationer i.
            lo, hi = max(month_start(min(dates)), fy.start), min(month_end(max(dates)), fy.end)
        else:
            lo = hi = fy.start
        covered = frozenset(months_between(month_start(lo), hi)) if dates else frozenset()
        opening = {a: v for a, v in openings.get(start, {}).items() if v != 0}
        status = "known" if start in openings else "missing"
        ledger.years.append(
            YearData(
                fiscal_year=fy,
                vouchers=vouchers,
                opening=opening,
                has_vouchers=True,
                opening_status=status,
                covered_months=covered,
            )
        )
        ranges[start] = (lo, hi)
        if dates and covered != frozenset(fy.months()):
            issues.append(
                ParseIssue(
                    None,
                    "PARTIAL_YEAR",
                    f"Räkenskapsåret {fy.label}: filen täcker {lo.isoformat()}–{hi.isoformat()}. "
                    "Övriga månader räknas som saknade, inte som noll.",
                    Severity.INFO,
                )
            )
    ledger.sort()
    for year in ledger.years:
        if year.opening_status == "missing":
            issues.append(
                ParseIssue(
                    None,
                    "MISSING_OPENING_BALANCES",
                    f"Ingående balanser saknas för räkenskapsåret {year.fiscal_year.label}. Balansposter (kassa, "
                    "soliditet, likviditet) visas som otillräckligt underlag tills IB finns – t.ex. från "
                    "föregående års SIE-fil eller IB-rader i filen (verifikation ”IB” eller text ”Ingående balans”).",
                )
            )
    if not fiscal_years:
        issues.append(
            ParseIssue(
                None,
                "ASSUMED_FISCAL_YEAR",
                "Filen anger inte räkenskapsår; "
                + (
                    "kalenderår antas."
                    if not fiscal_year_start_month or fiscal_year_start_month == 1
                    else f"räkenskapsår med start i månad {fiscal_year_start_month} antas."
                ),
                Severity.INFO,
            )
        )
    return TabularResult(ledger, issues, columns, fmt, encoding, delimiter, ranges)


def parse_csv(
    raw: bytes, *, fiscal_years: Sequence[FiscalYear] = (), fiscal_year_start_month: int | None = None
) -> TabularResult:
    rows, encoding, delimiter = read_csv_rows(raw)
    return parse_rows(
        rows,
        fmt="CSV",
        encoding=encoding,
        delimiter=delimiter,
        fiscal_years=fiscal_years,
        fiscal_year_start_month=fiscal_year_start_month,
    )


def parse_xlsx(
    raw: bytes, *, fiscal_years: Sequence[FiscalYear] = (), fiscal_year_start_month: int | None = None
) -> TabularResult:
    sheets = read_xlsx_rows(raw)
    for title, rows in sheets:
        if find_header(rows) is not None:
            return parse_rows(
                rows,
                fmt="XLSX",
                sheet=title,
                fiscal_years=fiscal_years,
                fiscal_year_start_month=fiscal_year_start_month,
            )
    names = ", ".join(title for title, _ in sheets) or "–"
    raise TabularFormatError(
        f"Hittar ingen rubrikrad med verifikation, datum, konto och belopp i något blad (blad: {names})."
    )
